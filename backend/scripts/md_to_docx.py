"""Convert reports/final_report.md to reports/final_report.docx.

    python scripts/md_to_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches

B = Path(__file__).resolve().parents[1]
REPORT = B / "reports" / "final_report.md"
OUT = B / "reports" / "final_report.docx"


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    return text.strip()


def add_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    first_heading = True
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 1
            if i < len(lines) and set(
                lines[i].replace("|", "").replace("-", "").replace(":", "").strip()
            ) == set():
                i += 1
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for j, value in enumerate(header):
                table.rows[0].cells[j].text = strip_md(value)
            for row in rows:
                cells = table.add_row().cells
                for j, value in enumerate(row[: len(header)]):
                    cells[j].text = strip_md(value)
            doc.add_paragraph()
            continue

        if line.startswith("!["):
            match = re.search(r"\((.+?)\)", line)
            if match:
                image = B / "reports" / match.group(1)
                if image.exists():
                    doc.add_picture(str(image), width=Inches(6.2))
            i += 1
            continue

        if line.startswith("# "):
            if not first_heading:
                doc.add_page_break()
            doc.add_heading(strip_md(line[2:]), level=1)
            first_heading = False
        elif line.startswith("## "):
            doc.add_heading(strip_md(line[3:]), level=2)
        elif line.startswith("### "):
            doc.add_heading(strip_md(line[4:]), level=3)
        elif line.startswith("- "):
            doc.add_paragraph(strip_md(line[2:]), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(strip_md(re.sub(r"^\d+\.\s", "", line)), style="List Number")
        elif line.strip():
            doc.add_paragraph(strip_md(line))
        i += 1


doc = Document()
add_markdown(doc, REPORT.read_text(encoding="utf-8"))
doc.save(OUT)
print(f"wrote {OUT}")
